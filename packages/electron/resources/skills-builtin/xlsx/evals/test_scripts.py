"""CLI regressions: kowork-python -B -m unittest discover -s evals -p 'test_*.py'."""
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"


class SheetPositionsTest(unittest.TestCase):
    def test_one_based_insertion_and_movement_preserve_active_sheet(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "source.xlsx"
            wb = Workbook()
            wb.active.title = "Original"
            wb.create_sheet("Last")
            wb.save(source)
            wb.close()

            for command, options, expected in (
                ("add", ["--name", "New", "--index", "1"], ["New", "Original", "Last"]),
                ("add", ["--name", "New", "--index", "3"], ["Original", "Last", "New"]),
                ("move", ["--sheet", "2", "--to-index", "1"], ["Last", "Original"]),
                ("move", ["--sheet", "1", "--to-index", "2"], ["Last", "Original"]),
                ("add", ["--name", "New", "--index", "0"], None),
                ("move", ["--sheet", "1", "--to-index", "0"], None),
                ("add", ["--name", "New", "--index", "4"], None),
                ("move", ["--sheet", "1", "--to-index", "3"], None),
            ):
                with self.subTest(command=command, options=options):
                    output = root / "output.xlsx"
                    output.unlink(missing_ok=True)
                    result = subprocess.run(
                        [sys.executable, "-B", str(SCRIPTS / "sheets.py"), command,
                         str(source), "-o", str(output), *options],
                        capture_output=True, text=True,
                    )
                    self.assertEqual(result.returncode, 0 if expected else 1, result.stderr)
                    if expected:
                        actual = load_workbook(output)
                        self.assertEqual(actual.sheetnames, expected)
                        self.assertEqual(actual.active.title, "Original")
                        actual.close()
                    else:
                        self.assertFalse(output.exists())
